import json
import random
import re
from datetime import date, datetime
from pathlib import Path
from urllib.parse import quote

from openpyxl import Workbook, load_workbook
from playwright.sync_api import TimeoutError as PlaywrightTimeoutError
from playwright.sync_api import sync_playwright


BASE_DIR = Path(__file__).resolve().parent
CONTACTS_FILE = BASE_DIR / "contacts.xlsx"
PROFILE_DIR = BASE_DIR / "whatsapp_profile"
SCREENSHOT_DIR = BASE_DIR / "screenshots"
REPORT_DATE = date.today().isoformat()


def pause(page):
    """Wait a human-like interval between browser actions."""
    page.wait_for_timeout(random.randint(2000, 5000))


def first_visible(page, selectors):
    """Return the first visible locator matching any selector."""
    for selector in selectors:
        locator = page.locator(selector).first
        try:
            if locator.count() and locator.is_visible():
                return locator
        except Exception:
            continue
    return None


def phone_digits(value):
    """Convert an Excel phone value to digits for WhatsApp's URL."""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"\D", "", str(value or ""))


def read_contacts():
    if not CONTACTS_FILE.exists():
        raise FileNotFoundError(f"Could not find {CONTACTS_FILE}")

    sheet = load_workbook(CONTACTS_FILE, data_only=True).active
    headers = {
        str(cell.value).strip().lower(): index
        for index, cell in enumerate(sheet[1])
        if cell.value is not None
    }

    if "name" not in headers or "phone" not in headers:
        raise ValueError("contacts.xlsx must have Name and Phone columns.")

    contacts = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        name = str(row[headers["name"]] or "").strip()
        phone = row[headers["phone"]]
        message = (
            str(row[headers["message"]] or "").strip()
            if "message" in headers
            else ""
        )

        # Ignore fully blank spreadsheet rows.
        if not name and not phone and not message:
            continue

        contacts.append({
            "name": name,
            "phone": str(phone or "").strip(),
            "message_template": message or "Hello {name}, I hope you're doing well.",
        })

    return contacts


def find_search_box(page):
    return first_visible(page, [
        '[aria-label="Search input textbox"]',
        'div[contenteditable="true"][data-tab="3"]',
        'input[placeholder*="Search"]',
    ])


def find_message_box(page):
    return first_visible(page, [
        'footer [contenteditable="true"][role="textbox"]',
        'footer [contenteditable="true"]',
        '[aria-label="Type a message"]',
        'div[contenteditable="true"][data-tab="10"]',
    ])


def open_contact(page, name, digits):
    """Try WhatsApp search, then fall back to opening the phone-number chat."""
    search_box = find_search_box(page)

    if search_box:
        for query in (name, digits):
            if not query:
                continue
            try:
                search_box.click()
                search_box.fill(query)
                pause(page)

                pane = page.locator("#pane-side")
                results = pane if pane.count() else page
                match = results.get_by_text(query, exact=True).last

                if match.count() and match.is_visible():
                    match.click(timeout=5000)
                    page.wait_for_selector("#main", timeout=10000)
                    page.wait_for_selector(
                        'footer [contenteditable="true"]', timeout=12000
                    )
                    if find_message_box(page):
                        return True
            except PlaywrightTimeoutError:
                pass
            except Exception:
                # Search results vary across WhatsApp Web versions.
                pass

    if not digits:
        return False

    # The URL fallback also supports numbers that are not saved as contacts.
    page.goto(
        f"https://web.whatsapp.com/send?phone={quote(digits)}",
        wait_until="domcontentloaded",
    )
    try:
        page.wait_for_selector(
            'footer [contenteditable="true"]', timeout=15000
        )
        return find_message_box(page) is not None
    except PlaywrightTimeoutError:
        return False


def extract_last_incoming_messages(page):
    """Return up to three latest text messages received in the open chat."""
    page.wait_for_timeout(1000)
    messages = page.locator("div.message-in .selectable-text").all_inner_texts()
    return [text.strip() for text in messages if text.strip()][-3:]


def save_reports(records, generated_at):
    json_path = BASE_DIR / f"whatsapp_report_{REPORT_DATE}.json"
    excel_path = BASE_DIR / f"whatsapp_report_{REPORT_DATE}.xlsx"

    report = {
        "title": "Social Eagle · CAIE Architect Program",
        "generated_at": generated_at,
        "contacts_file": CONTACTS_FILE.name,
        "results": records,
    }
    json_path.write_text(
        json.dumps(report, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Summary"
    sheet.append(["Social Eagle · CAIE Architect Program"])
    sheet.append(["Generated at", generated_at])
    sheet.append([])
    sheet.append([
        "Name",
        "Phone",
        "Status",
        "Message",
        "Screenshot",
        "Last 3 incoming messages",
        "Error",
    ])

    for item in records:
        sheet.append([
            item.get("name", ""),
            item.get("phone", ""),
            item.get("status", ""),
            item.get("message_sent", ""),
            item.get("screenshot", ""),
            "\n".join(item.get("last_3_incoming_messages", [])),
            item.get("error", ""),
        ])

    sheet.freeze_panes = "A5"
    for column, width in {
        "A": 22, "B": 20, "C": 20, "D": 45,
        "E": 40, "F": 60, "G": 45,
    }.items():
        sheet.column_dimensions[column].width = width

    workbook.save(excel_path)
    print(f"Reports saved:\n  {json_path}\n  {excel_path}")


def main():
    contacts = read_contacts()
    generated_at = datetime.now().astimezone().isoformat(timespec="seconds")
    records = []
    SCREENSHOT_DIR.mkdir(parents=True, exist_ok=True)

    with sync_playwright() as playwright:
        context = playwright.chromium.launch_persistent_context(
            user_data_dir=str(PROFILE_DIR),
            headless=False,
        )
        page = context.pages[0] if context.pages else context.new_page()

        try:
            page.goto("https://web.whatsapp.com", wait_until="domcontentloaded")

            # The persistent profile remembers the login after the first run.
            try:
                page.wait_for_selector("#pane-side", timeout=5000)
            except PlaywrightTimeoutError:
                input(
                    "Scan the WhatsApp Web QR code in the browser, "
                    "then press Enter here..."
                )
                page.wait_for_selector("#pane-side", timeout=180000)

            for contact in contacts:
                name = contact["name"]
                digits = phone_digits(contact["phone"])
                message = contact["message_template"].replace("{name}", name)

                result = {
                    "name": name,
                    "phone": contact["phone"],
                    "message_sent": message,
                    "status": "not_started",
                    "screenshot": "",
                    "last_3_incoming_messages": [],
                    "error": "",
                }

                try:
                    if not name or not digits:
                        result["status"] = "invalid_contact"
                        result["error"] = "Name or phone number is missing."
                        records.append(result)
                        continue

                    if not open_contact(page, name, digits):
                        result["status"] = "contact_not_found"
                        result["error"] = (
                            "Could not find an open chat or message box."
                        )
                        records.append(result)
                        pause(page)
                        continue

                    pause(page)
                    message_box = find_message_box(page)
                    if not message_box:
                        result["status"] = "send_failed"
                        result["error"] = "Could not locate the message box."
                        records.append(result)
                        pause(page)
                        continue

                    message_box.fill(message)
                    pause(page)
                    message_box.press("Enter")
                    result["status"] = "sent_unconfirmed"

                    # A visible outgoing bubble confirms the message appeared
                    # in the chat; this is a UI confirmation, not a delivery receipt.
                    outgoing = (
                        page.locator("div.message-out")
                        .filter(has_text=message)
                        .last
                    )
                    outgoing.wait_for(state="visible", timeout=15000)
                    result["status"] = "sent"
                    pause(page)

                    safe_name = re.sub(r"[^A-Za-z0-9_-]+", "_", name).strip("_")
                    screenshot_path = (
                        SCREENSHOT_DIR
                        / f"{REPORT_DATE}_{safe_name or 'contact'}.png"
                    )
                    try:
                        outgoing.screenshot(path=str(screenshot_path))
                        result["screenshot"] = str(
                            screenshot_path.relative_to(BASE_DIR)
                        )
                    except Exception as exc:
                        result["screenshot_error"] = str(exc)

                    try:
                        result["last_3_incoming_messages"] = (
                            extract_last_incoming_messages(page)
                        )
                    except Exception as exc:
                        result["extraction_error"] = str(exc)

                except PlaywrightTimeoutError as exc:
                    if result["status"] == "sent_unconfirmed":
                        result["error"] = (
                            "Message was submitted, but an outgoing bubble "
                            f"was not confirmed: {exc}"
                        )
                    else:
                        result["status"] = "send_failed"
                        result["error"] = f"Timed out: {exc}"
                except Exception as exc:
                    result["status"] = (
                        "send_failed"
                        if result["status"] == "not_started"
                        else result["status"]
                    )
                    result["error"] = str(exc)

                records.append(result)
                pause(page)

        finally:
            # Save partial results even if the run encounters an error.
            save_reports(records, generated_at)
            context.close()


if __name__ == "__main__":
    main()