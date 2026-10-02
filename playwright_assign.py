"""
WhatsApp Message Sender + Smart Data Extractor
==============================================
A Playwright (Python) bot that drives WhatsApp Web the way a person would:

  1. Logs in to WhatsApp Web. Scan the QR code on the first run; the session is
     saved in ./whatsapp_session so later runs skip the QR step.
  2. Reads contacts from contacts.xlsx  (columns: Name, Phone, Message).
  3. For each contact: searches by number (then by name), types a personalised
     message ({name} -> real name), waits for the sent tick and screenshots it.
  4. Re-opens every contact it reached and extracts their last 3 messages.
  5. Writes whatsapp_report_YYYY-MM-DD.json (full details) and
     whatsapp_report_YYYY-MM-DD.xlsx (summary).

Setup:   pip install playwright openpyxl  &&  playwright install chromium
Run:     python playwright_assign.py
"""

import json
import logging
import random
import re
import sys
from collections import Counter
from contextlib import suppress
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from playwright.sync_api import TimeoutError as PWTimeout
from playwright.sync_api import sync_playwright

# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------
BASE_DIR = Path(__file__).resolve().parent
CONTACTS_FILE = BASE_DIR / "contacts.xlsx"
SESSION_DIR = BASE_DIR / "whatsapp_session"   # browser profile that keeps you logged in
SCREENSHOT_DIR = BASE_DIR / "screenshots"
REPORT_DIR = BASE_DIR
LOG_FILE = BASE_DIR / "whatsapp_bot.log"
WHATSAPP_URL = "https://web.whatsapp.com"

DEFAULT_TEMPLATE = "Hi {name}, hope you're doing well! Just checking in from our team."
DEFAULT_COUNTRY_CODE = "91"       # added to 10-digit numbers written without a country code
MIN_DELAY, MAX_DELAY = 2.0, 5.0   # seconds of random "human" pause between actions
TYPING_DELAY_MS = (40, 120)       # per-keystroke delay range while typing
LOGIN_TIMEOUT_MS = 180_000        # time allowed to scan the QR code
SEND_CONFIRM_TIMEOUT_MS = 30_000  # time allowed for the sent tick to appear
MESSAGES_TO_EXTRACT = 3
DRY_RUN = False                   # True = open chats and extract messages, but never send

# WhatsApp Web changes its markup often, so every selector lives here: a broken
# one is a one-line fix. Comma-separated entries are fallbacks for other layouts.
SEL = {
    "chat_list": '#pane-side, div[aria-label="Chat list"]',
    "qr_code": 'canvas[aria-label*="scan" i], div[data-ref] canvas',
    "search_box": '#side div[contenteditable="true"], #side input[type="text"], '
                  'div[contenteditable="true"][data-tab="3"], input[aria-label*="Search" i]',
    "search_result": '#pane-side [role="listitem"], #pane-side [role="row"]',
    "message_box": '#main footer div[contenteditable="true"], '
                   'div[contenteditable="true"][aria-label^="Type a message"]',
    "send_button": '#main footer button[aria-label="Send"], #main footer span[data-icon="send"]',
    "chat_header": "#main header",
    "chat_panel": "#main",
    "outgoing": "#main div.message-out",
    "incoming": "#main div.message-in",
    "sent_tick": 'span[data-icon*="check"], span[aria-label*="Sent"], span[aria-label*="Delivered"], '
                 'span[aria-label=" Read "], span[aria-label="Read"]',
}
NO_RESULTS_TEXT = re.compile(r"no (chats|contacts|results)\b.*found", re.I)
INVALID_NUMBER_TEXT = re.compile(r"phone number shared via url is invalid", re.I)

# Result statuses
SENT, PENDING, DRY = "Sent", "Pending", "Dry Run"
NOT_FOUND, FAILED, SKIPPED, NOT_PROCESSED = "Not Found", "Failed", "Skipped", "Not Processed"
ALL_STATUSES = (SENT, PENDING, DRY, NOT_FOUND, FAILED, SKIPPED, NOT_PROCESSED)
REACHED = (SENT, PENDING, DRY)    # contacts whose chat was opened in phase 1
STATUS_COLORS = {SENT: "C6EFCE", DRY: "DDEBF7", PENDING: "FFEB9C", SKIPPED: "FFEB9C",
                 NOT_FOUND: "FFC7CE", FAILED: "FFC7CE", NOT_PROCESSED: "EDEDED"}

log = logging.getLogger("whatsapp_bot")


# ---------------------------------------------------------------------------
# Small helpers
# ---------------------------------------------------------------------------
def setup_logging():
    with suppress(Exception):
        sys.stderr.reconfigure(errors="replace")   # don't choke on emoji in names
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s | %(levelname)-7s | %(message)s",
        datefmt="%H:%M:%S",
        handlers=[logging.StreamHandler(), logging.FileHandler(LOG_FILE, encoding="utf-8")],
    )


def human_pause(page, min_s=MIN_DELAY, max_s=MAX_DELAY):
    """Random pause so the bot never acts at machine speed (reduces ban risk)."""
    page.wait_for_timeout(random.uniform(min_s, max_s) * 1000)


def wait_until(page, condition, timeout_ms, step_ms=500):
    """Poll condition() until it is truthy or timeout_ms has passed."""
    for _ in range(max(1, timeout_ms // step_ms)):
        if condition():
            return True
        page.wait_for_timeout(step_ms)
    return bool(condition())


def now_iso():
    return datetime.now().isoformat(timespec="seconds")


def short_error(exc):
    """First line of an exception (Playwright errors carry long call logs)."""
    first = (str(exc).strip().splitlines() or [""])[0]
    return f"{type(exc).__name__}: {first}"[:250]


def rel(path):
    try:
        return path.relative_to(BASE_DIR).as_posix()
    except ValueError:
        return path.as_posix()


def clean(value):
    """Strip characters Excel cannot store."""
    return ILLEGAL_CHARACTERS_RE.sub("", value) if isinstance(value, str) else value


# ---------------------------------------------------------------------------
# Contacts (Excel input)
# ---------------------------------------------------------------------------
def normalize_phone(raw):
    """'+91 98765-43210' / 919876543210 / 9876543210.0  ->  '919876543210' ('' if invalid)."""
    if raw is None:
        return ""
    if isinstance(raw, float):
        raw = int(raw)
    digits = re.sub(r"\D", "", str(raw))
    if len(digits) == 10 and DEFAULT_COUNTRY_CODE:
        digits = DEFAULT_COUNTRY_CODE + digits
    return digits if 11 <= len(digits) <= 15 else ""


def personalise(template, name):
    """Replace {name} (any case) in the template; blank template -> DEFAULT_TEMPLATE."""
    return re.sub(r"\{name\}", lambda _: name or "there", template or DEFAULT_TEMPLATE, flags=re.I).strip()


def read_contacts(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    rows = wb.active.iter_rows(values_only=True)
    header = [str(h or "").strip().lower() for h in next(rows, ())]

    def column(*keys):
        return next((i for i, h in enumerate(header) if any(k in h for k in keys)), None)

    name_col, phone_col, msg_col = column("name"), column("phone", "mobile", "number"), column("message", "template")
    if name_col is None or phone_col is None:
        wb.close()
        raise SystemExit(f"{path.name} needs 'Name' and 'Phone' columns (found: {header})")

    contacts = []
    for row_no, row in enumerate(rows, start=2):
        def cell(i):
            return row[i] if i is not None and i < len(row) else None

        name, raw_phone = str(cell(name_col) or "").strip(), cell(phone_col)
        if not name and raw_phone in (None, ""):
            continue                                  # blank row
        if isinstance(raw_phone, float):
            raw_phone = int(raw_phone)                # Excel stored the number as 9.19e11
        contacts.append({
            "row": row_no,
            "name": name,
            "phone_raw": str(raw_phone or "").strip(),
            "phone": normalize_phone(raw_phone),
            "template": str(cell(msg_col) or "").strip(),
        })
    wb.close()
    return contacts


def create_sample_contacts(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Contacts"
    ws.append(["Name", "Phone", "Message"])
    ws.append(["Ravi", "+91XXXXXXXXXX", "Hi {name}, your order is ready for pickup!"])
    ws.append(["Priya", "+91XXXXXXXXXX", ""])   # blank Message -> DEFAULT_TEMPLATE
    for col, width in zip("ABC", (20, 18, 60)):
        ws.column_dimensions[col].width = width
        ws[f"{col}1"].font = Font(bold=True)
    wb.save(path)


def new_result(contact):
    return {
        "row": contact["row"],
        "name": contact["name"],
        "phone": f"+{contact['phone']}" if contact["phone"] else contact["phone_raw"],
        "message": personalise(contact["template"], contact["name"]),
        "status": NOT_PROCESSED,
        "found_via": None,
        "delivery": None,
        "sent_at": None,
        "screenshot": None,
        "error": None,
        "last_messages": [],
        "extraction": None,
    }


# ---------------------------------------------------------------------------
# WhatsApp Web actions
# ---------------------------------------------------------------------------
def login(page):
    log.info("Opening WhatsApp Web ...")
    page.goto(WHATSAPP_URL, wait_until="domcontentloaded")
    # A saved session shows the chat list; a new one shows the QR code first.
    page.wait_for_selector(f'{SEL["chat_list"]}, {SEL["qr_code"]}', timeout=90_000)
    if not page.locator(SEL["chat_list"]).first.is_visible():
        log.info("Scan the QR code: phone > WhatsApp > Linked devices > Link a device "
                 "(waiting up to %d s)", LOGIN_TIMEOUT_MS // 1000)
        page.wait_for_selector(SEL["chat_list"], timeout=LOGIN_TIMEOUT_MS)
    log.info("Logged in - letting chats finish loading")
    page.wait_for_timeout(5000)


def chat_matches(page, contact):
    """True if the open chat's header shows exactly this contact's name or number.
    This guard is what stops the bot from ever messaging the wrong person."""
    header = page.locator(SEL["chat_header"]).first
    try:
        if not header.count():
            return False
        lines = [ln.strip() for ln in header.inner_text(timeout=2000).splitlines() if ln.strip()]
    except Exception:
        return False
    name, tail = contact["name"].casefold(), contact["phone"][-10:]
    return any((name and ln.casefold() == name) or (tail and re.sub(r"\D", "", ln).endswith(tail))
               for ln in lines)


def wait_for_chat(page, contact, timeout_ms=8000):
    return wait_until(page, lambda: page.locator(SEL["message_box"]).count() > 0
                      and chat_matches(page, contact), timeout_ms)


def search_and_open(page, query, contact):
    """Type `query` into the sidebar search and open the matching chat."""
    log.info("   searching for '%s'", query)
    page.wait_for_selector(SEL["search_box"], timeout=15_000)
    box = page.locator(SEL["search_box"]).first
    box.click()
    box.press("Control+A")
    box.press("Backspace")
    box.press_sequentially(query, delay=random.randint(*TYPING_DELAY_MS))
    human_pause(page)                                  # also lets WhatsApp filter the results

    if page.get_by_text(NO_RESULTS_TEXT).count():
        log.info("   no results for '%s'", query)
        return False

    # Prefer a result titled exactly like the contact, otherwise the top result.
    exact = page.locator("#pane-side").get_by_title(contact["name"], exact=True) if contact["name"] else None
    if exact is not None and exact.count():
        exact.first.click()
    else:
        box.press("Enter")                             # Enter opens the top search result
    if wait_for_chat(page, contact):
        return True

    results = page.locator(SEL["search_result"])
    if results.count():                                # Enter didn't do it - click the top result
        results.first.click()
        if wait_for_chat(page, contact):
            return True
    log.info("   '%s' did not open a chat for %s", query, contact["name"] or contact["phone"])
    return False


def open_via_link(page, phone):
    """WhatsApp's click-to-chat link - reaches numbers that aren't saved in your phone."""
    log.info("   opening chat via direct link for +%s", phone)
    page.goto(f"{WHATSAPP_URL}/send?phone={phone}", wait_until="domcontentloaded")
    box = page.locator(SEL["message_box"])
    invalid = page.get_by_text(INVALID_NUMBER_TEXT)
    box.or_(invalid).first.wait_for(timeout=60_000)    # either the chat or the "invalid" popup
    if invalid.count():
        log.info("   +%s is not on WhatsApp", phone)
        ok = page.get_by_role("button", name=re.compile(r"^ok$", re.I))
        if ok.count():
            ok.first.click()
        return False
    return True


def open_chat(page, contact):
    """Search by number, then by name, then fall back to the direct link.
    Returns how the chat was found, or None if the contact can't be reached."""
    if page.locator(SEL["message_box"]).count() and chat_matches(page, contact):
        return "already open"
    for how, query in (("search: number", contact["phone"]), ("search: name", contact["name"])):
        if query and search_and_open(page, query, contact):
            return how
    if contact["phone"] and open_via_link(page, contact["phone"]):
        return "direct link"
    return None


def last_outgoing_id(page):
    """Id of the newest sent bubble - changes the moment a new message lands in the chat."""
    return page.evaluate(
        """(sel) => {
            const all = document.querySelectorAll(sel);
            if (!all.length) return null;
            const row = all[all.length - 1].closest('[data-id]');
            return row ? row.dataset.id : String(all.length);
        }""",
        SEL["outgoing"],
    )


def delivery_status(bubble):
    labels = bubble.locator(SEL["sent_tick"]).evaluate_all(
        "els => els.map(e => (e.getAttribute('aria-label') || '') + ' ' + (e.getAttribute('data-icon') || '')).join(' ')"
    ).lower()
    if "read" in labels:
        return "Read"
    if "deliver" in labels or "dblcheck" in labels:
        return "Delivered"
    return "Sent"


def send_message(page, text):
    """Type into the 'Type a message' box, send, and wait for the tick.
    Returns (delivery status, locator of the sent bubble)."""
    page.wait_for_selector(SEL["message_box"], timeout=15_000)
    box = page.locator(SEL["message_box"]).first
    box.click()
    box.press("Control+A")
    box.press("Backspace")                             # clear any leftover draft

    lines = text.split("\n")
    for i, line in enumerate(lines):
        if line:
            box.press_sequentially(line, delay=random.randint(*TYPING_DELAY_MS))
        if i < len(lines) - 1:
            box.press("Shift+Enter")                   # new line without sending
    if not box.evaluate("el => el.innerText.trim().length > 0 || !!el.querySelector('img')"):
        log.info("   keystrokes didn't land - pasting the text instead")
        box.click()
        page.keyboard.insert_text(text)
    human_pause(page)                                  # "re-read" before hitting send

    before = last_outgoing_id(page)
    box.press("Enter")
    if not wait_until(page, lambda: last_outgoing_id(page) != before, 10_000):
        log.info("   Enter didn't send - clicking the send button")
        page.locator(SEL["send_button"]).first.click(timeout=5000)
        if not wait_until(page, lambda: last_outgoing_id(page) != before, 10_000):
            raise RuntimeError("sent message never appeared in the chat")

    bubble = page.locator(SEL["outgoing"]).last
    try:
        # "attached", not "visible": the tick is a tiny icon that can report as hidden
        bubble.locator(SEL["sent_tick"]).first.wait_for(state="attached", timeout=SEND_CONFIRM_TIMEOUT_MS)
    except PWTimeout:
        return "Pending", bubble                       # still showing the clock icon
    return delivery_status(bubble), bubble


def take_screenshot(page, contact, bubble=None, prefix="sent"):
    """Screenshot the chat panel (full page for errors) and return the file path."""
    try:
        folder = SCREENSHOT_DIR / datetime.now().strftime("%Y-%m-%d")
        folder.mkdir(parents=True, exist_ok=True)
        safe = re.sub(r"[^\w-]+", "_", contact["name"] or contact["phone"] or "contact").strip("_")
        path = folder / f"{prefix}_row{contact['row']}_{safe}_{datetime.now():%H%M%S}.png"
        if bubble is not None:
            bubble.scroll_into_view_if_needed(timeout=5000)
        panel = page.locator(SEL["chat_panel"])
        (panel if prefix != "error" and panel.count() else page).screenshot(path=str(path))
        return rel(path)
    except Exception as exc:
        log.warning("   could not save screenshot: %s", short_error(exc))
        return None


# ---------------------------------------------------------------------------
# Smart data extraction
# ---------------------------------------------------------------------------
# Rebuilds each incoming message's text (emoji are <img alt="..."> in WhatsApp's
# DOM, so plain innerText would drop them) and reads the hidden metadata
# attribute WhatsApp stores on every message: "[10:32 am, 30/09/2026] Ravi: ".
EXTRACT_JS = r"""
els => els.map(el => {
    const meta = el.querySelector('[data-pre-plain-text]');
    const body = el.querySelector('span.selectable-text');
    const walk = n => n.nodeType === Node.TEXT_NODE ? n.textContent
                    : n.nodeName === 'IMG' ? (n.alt || '')
                    : n.nodeName === 'BR' ? '\n'
                    : Array.from(n.childNodes).map(walk).join('');
    return {
        meta: meta ? meta.getAttribute('data-pre-plain-text') : '',
        text: body ? walk(body).trim() : '',
        raw: el.innerText || '',
    };
})
"""
META_RE = re.compile(r"^\[(?P<time>[^,\]]+),\s*(?P<date>[^\]]+)\]\s*(?P<sender>.*?):\s*$")
TIME_RE = re.compile(r"\b\d{1,2}:\d{2}(?:\s?[ap]\.?m\.?)?", re.I)


def scroll_chat_up(page):
    box = page.locator(SEL["chat_panel"]).bounding_box()
    if box:
        page.mouse.move(box["x"] + box["width"] / 2, box["y"] + box["height"] / 2)
        page.mouse.wheel(0, -3000)
        page.wait_for_timeout(1500)


def extract_last_messages(page, contact, count=MESSAGES_TO_EXTRACT):
    """Last `count` messages *from the contact* in the open chat, oldest first."""
    incoming = page.locator(SEL["incoming"])
    for _ in range(3):                                 # older messages load lazily on scroll
        if incoming.count() >= count:
            break
        scroll_chat_up(page)

    messages = []
    for item in incoming.evaluate_all(EXTRACT_JS)[-count:]:
        meta = META_RE.match(item["meta"].strip())
        time = TIME_RE.search(item["raw"])
        messages.append({
            "sender": meta["sender"] if meta else (contact["name"] or f"+{contact['phone']}"),
            "date": meta["date"].strip() if meta else "",
            "time": meta["time"].strip() if meta else (time.group(0) if time else ""),
            "type": "text" if item["text"] else "media/other",
            "text": item["text"] or "[non-text message: photo, sticker, voice note, etc.]",
        })
    return messages


def format_msg(m):
    stamp = " ".join(p for p in (m["date"], m["time"]) if p)
    return f"[{stamp}] {m['text']}" if stamp else m["text"]


# ---------------------------------------------------------------------------
# The two phases
# ---------------------------------------------------------------------------
def send_to_contact(page, contact, result):
    if not contact["phone"] and not contact["name"]:
        result.update(status=SKIPPED, error=f"No valid phone number or name (phone: '{contact['phone_raw']}')")
        return
    bubble = None
    try:
        found_via = open_chat(page, contact)
        if not found_via:
            reason = "Contact not found on WhatsApp"
            if not contact["phone"]:
                reason += f" (no valid phone number to fall back on: '{contact['phone_raw']}')"
            result.update(status=NOT_FOUND, error=reason)
            return
        result["found_via"] = found_via
        human_pause(page)
        if DRY_RUN:
            result["status"] = DRY
        else:
            delivery, bubble = send_message(page, result["message"])
            result.update(status=PENDING if delivery == "Pending" else SENT,
                          delivery=delivery, sent_at=now_iso())
        result["screenshot"] = take_screenshot(page, contact, bubble, prefix="dryrun" if DRY_RUN else "sent")
    except Exception as exc:                           # one bad contact must not crash the run
        log.error("   error: %s", short_error(exc))
        result.update(status=FAILED, error=short_error(exc))
        result["screenshot"] = take_screenshot(page, contact, prefix="error")


def extract_from_contact(page, contact, result):
    try:
        if not open_chat(page, contact):
            result["extraction"] = "Chat not found"
            return
        human_pause(page)
        result["last_messages"] = extract_last_messages(page, contact)
        n = len(result["last_messages"])
        result["extraction"] = f"{n} message(s) extracted" if n else "No messages from contact yet"
    except Exception as exc:
        log.error("   extraction error: %s", short_error(exc))
        result["extraction"] = f"Failed: {short_error(exc)}"


def run_bot(page, contacts, results):
    login(page)
    total = len(contacts)

    log.info("=== Phase 1: sending messages to %d contact(s)%s ===", total, " [DRY RUN]" if DRY_RUN else "")
    for i, (contact, result) in enumerate(zip(contacts, results), 1):
        if page.is_closed():
            raise RuntimeError("the browser window was closed")
        log.info("[%d/%d] %s  %s", i, total, result["name"] or "-", result["phone"] or "-")
        send_to_contact(page, contact, result)
        log.info("   -> %s%s", result["status"], f" ({result['delivery']})" if result["delivery"] else "")
        if i < total:
            human_pause(page)

    reached = [(c, r) for c, r in zip(contacts, results) if r["status"] in REACHED]
    log.info("=== Phase 2: extracting last %d messages from %d contact(s) ===", MESSAGES_TO_EXTRACT, len(reached))
    for i, (contact, result) in enumerate(reached, 1):
        if page.is_closed():
            raise RuntimeError("the browser window was closed")
        log.info("[%d/%d] %s", i, len(reached), result["name"] or result["phone"])
        extract_from_contact(page, contact, result)
        log.info("   -> %s", result["extraction"])


# ---------------------------------------------------------------------------
# Reports
# ---------------------------------------------------------------------------
def write_excel(report, path):
    n = MESSAGES_TO_EXTRACT
    columns = [("#", 5), ("Name", 20), ("Phone", 16), ("Status", 12), ("Delivery", 11),
               ("Found Via", 15), ("Sent At", 20), ("Message", 45)]
    columns += [(f"Last Msg {i}" + (" (newest)" if i == 1 else ""), 38) for i in range(1, n + 1)]
    columns += [("Extraction", 24), ("Screenshot", 42), ("Error", 40)]

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append([c for c, _ in columns])
    for i, r in enumerate(report["results"], 1):
        latest = [format_msg(m) for m in reversed(r["last_messages"])] + [""] * n
        ws.append([clean(v) for v in (
            i, r["name"], r["phone"], r["status"], r["delivery"], r["found_via"], r["sent_at"],
            r["message"], *latest[:n], r["extraction"], r["screenshot"], r["error"])])
        ws.cell(ws.max_row, 4).fill = PatternFill("solid", fgColor=STATUS_COLORS.get(r["status"], "FFFFFF"))

    for idx, (_, width) in enumerate(columns, 1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="075E54")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    info = wb.create_sheet("Run Info")
    rows = [("Report date", report["report_date"]), ("Run started", report["run_started"]),
            ("Run finished", report["run_finished"]), ("Duration (s)", report["duration_seconds"]),
            ("Dry run", "Yes" if report["dry_run"] else "No"), ("Contacts file", report["contacts_file"])]
    rows += [(k.replace("_", " ").title(), v) for k, v in report["totals"].items()]
    for row in rows:
        info.append(row)
    for cell in info["A"]:
        cell.font = Font(bold=True)
    info.column_dimensions["A"].width = 18
    info.column_dimensions["B"].width = 24
    wb.save(path)


def save_reports(results, started):
    finished = datetime.now()
    stamp = started.strftime("%Y-%m-%d")
    counts = Counter(r["status"] for r in results)
    report = {
        "report_date": stamp,
        "run_started": started.isoformat(timespec="seconds"),
        "run_finished": finished.isoformat(timespec="seconds"),
        "duration_seconds": round((finished - started).total_seconds()),
        "dry_run": DRY_RUN,
        "contacts_file": CONTACTS_FILE.name,
        "totals": {"contacts": len(results),
                   **{s.lower().replace(" ", "_"): counts.get(s, 0) for s in ALL_STATUSES}},
        "results": results,
    }

    json_path = REPORT_DIR / f"whatsapp_report_{stamp}.json"
    json_path.write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")

    xlsx_path = REPORT_DIR / f"whatsapp_report_{stamp}.xlsx"
    try:
        write_excel(report, xlsx_path)
    except PermissionError:                            # the file is open in Excel
        xlsx_path = xlsx_path.with_name(f"{xlsx_path.stem}_{finished:%H%M%S}.xlsx")
        log.warning("Report is open in another program - saving as %s instead", xlsx_path.name)
        write_excel(report, xlsx_path)
    return report, json_path, xlsx_path


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------
def main():
    setup_logging()
    if not CONTACTS_FILE.exists():
        create_sample_contacts(CONTACTS_FILE)
        log.info("Created a sample %s - fill in your contacts and run again.", CONTACTS_FILE.name)
        return
    contacts = read_contacts(CONTACTS_FILE)
    if not contacts:
        log.error("No contacts found in %s", CONTACTS_FILE.name)
        return
    log.info("Loaded %d contact(s) from %s", len(contacts), CONTACTS_FILE.name)

    results = [new_result(c) for c in contacts]
    started = datetime.now()
    try:
        with sync_playwright() as p:
            context = p.chromium.launch_persistent_context(
                str(SESSION_DIR),
                headless=False,                        # WhatsApp needs a visible browser for the QR login
                locale="en-US",                        # English UI so the text-based checks match
                no_viewport=True,
                args=["--start-maximized"],
            )
            page = context.pages[0] if context.pages else context.new_page()
            try:
                run_bot(page, contacts, results)
            finally:
                with suppress(Exception):
                    context.close()
    except KeyboardInterrupt:
        log.warning("Stopped by user - saving what was done so far")
    except Exception as exc:
        log.error("Run stopped early: %s", short_error(exc))
    finally:
        report, json_path, xlsx_path = save_reports(results, started)
        totals = ", ".join(f"{k.replace('_', ' ')}: {v}" for k, v in report["totals"].items() if v)
        log.info("Done - %s", totals)
        log.info("Reports: %s | %s", json_path.name, xlsx_path.name)


if __name__ == "__main__":
    main()
